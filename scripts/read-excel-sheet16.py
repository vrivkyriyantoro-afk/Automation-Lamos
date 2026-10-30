import openpyxl

wb = openpyxl.load_workbook(r'fixtures\Analisa Data LAMOS SME Phase 2.xlsx', read_only=True, data_only=True)
sheet_names = wb.sheetnames

print("Total sheet:", len(sheet_names))
for idx, name in enumerate(sheet_names):
    print(f"Sheet {idx}: {name}")

target_sheet_name = None
for name in sheet_names:
    if "agunan" in name.lower() or "mck" in name.lower():
        target_sheet_name = name
        break

if not target_sheet_name and len(sheet_names) > 16:
    target_sheet_name = sheet_names[16]

print(f"\nMembaca sheet target: '{target_sheet_name}'")
sheet = wb[target_sheet_name]

rows = []
for i, row in enumerate(sheet.iter_rows(values_only=True)):
    if any(val is not None for val in row):
        # Bersihkan None
        cleaned_row = [str(val).strip() if val is not None else "" for val in row]
        # Hanya ambil kolom sampai kolom terakhir yang berisi
        while cleaned_row and not cleaned_row[-1]:
            cleaned_row.pop()
        if cleaned_row:
            rows.append(cleaned_row)
    if len(rows) >= 100:
        break

print(f"Total baris terbaca (sample 100): {len(rows)}")
for i, r in enumerate(rows[:25]):
    print(f"Row {i+1}: {r}")

