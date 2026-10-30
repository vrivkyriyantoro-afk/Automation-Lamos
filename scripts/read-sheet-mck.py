import openpyxl
import json

wb = openpyxl.load_workbook(r'fixtures\Analisa Data LAMOS SME Phase 2.xlsx', data_only=True)
sheet = wb['FEBucket-Agunan (MVP - McK)']

print("Membaca sheet: 'FEBucket-Agunan (MVP - McK)'")
print(f"Max row: {sheet.max_row}, Max column: {sheet.max_column}")

rows = []
for r in sheet.iter_rows(values_only=True):
    if any(val is not None for val in r):
        cleaned = [str(val).strip() if val is not None else "" for val in r]
        while cleaned and not cleaned[-1]:
            cleaned.pop()
        if cleaned:
            rows.append(cleaned)

print(f"Total baris berisi: {len(rows)}")

# Print 35 baris pertama untuk melihat struktur header & kolom
print("\n--- 35 BARIS PERTAMA ---")
for i, r in enumerate(rows[:35]):
    print(f"[{i+1}] {r}")

# Cari sections / headers
sections = []
for idx, r in enumerate(rows):
    if len(r) >= 2 and (not r[0] or r[0].lower().startswith('modul') or r[0].lower().startswith('section')):
        # Cek jika baris terlihat seperti nama bagian/sub-bagian
        pass

# Simpan dump text lengkap ke explore-output/sheet_mck_dump.txt agar bisa dianalisis mendalam
with open(r'explore-output\sheet_mck_dump.txt', 'w', encoding='utf-8') as f:
    for i, r in enumerate(rows):
        f.write(f"[{i+1}] " + " | ".join(r) + "\n")

print("\nDump berhasil disimpan di explore-output/sheet_mck_dump.txt")

